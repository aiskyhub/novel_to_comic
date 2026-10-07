"""Mechanical regression fixtures; these do not claim to validate fictional art quality."""
from __future__ import annotations

import copy
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import comic_pipeline as cp
import comic_sources as cs
import comic_layout as cl
from adaptation_fixtures import attach_adaptations


def fixture_script(source):
    script = cp.load_json(Path(cp.__file__).resolve().parents[1] / 'assets' / 'script-template.json')
    script.update(outline='测试夹具的全部事件顺序保留。', ending='以提供的最后一句结束。')
    script['style'].update(genre='测试', look='清晰线条', palette='灰色', selection_reason='自动化机械测试',
                           format='pages', width=1080, height=2400, font_size=54)
    script['characters'] = [{'id': 'char-a', 'name': '甲', 'aliases': [], 'importance': 'major',
        'design_tier': 'lead',
        'appearance': {'gender_presentation': 'masculine', 'requirements': '机械测试男性造型', 'source_unit_ids': []},
        'identity_card': {**{key: '机械身份锚点' for key in ('face', 'eyes_brows', 'nose_mouth', 'body', 'posture', 'temperament')}, 'invariants': ['测试脸型']},
        'appearance_versions': [{'id': 'base', 'description': '测试基础衣着', 'visual': {'costume': 'coat-a'}}],
        'comparison_with': [], 'distinctions': [],
        'narrative': {key: '测试人物' for key in ('goal', 'motivation', 'voice', 'arc')},
        'visual': {key: '测试设定' for key in ('face_shape', 'eyes', 'brows', 'nose_mouth', 'body', 'posture', 'hair', 'age')},
        'source_facts': [], 'design_notes': ['机械测试夹具，不代表实际人物图']}]
    script['settings'] = [{'id': 'room-a', 'description': '测试房间'}]
    body = [u for u in source['units'] if u['kind'] == 'body']
    for chapter in source['chapters']:
        if chapter['has_body']:
            script['scenes'].append({'id': 'scene-' + chapter['id'], 'chapter_id': chapter['id'], 'setting_id': 'room-a'})
    state = {'form': 'base', 'costume': 'coat-a', 'injuries': [], 'items': [], 'location': 'room-a', 'knowledge': []}
    for i, unit in enumerate(body, 1):
        event_id = f'ev{i}'
        pid = f'p{i}'
        script['events'].append({'id': event_id, 'description': unit['text'], 'source_unit_ids': [unit['id']]})
        script['panels'].append({
            'id': pid, 'chapter_id': unit['chapter_id'], 'scene_id': 'scene-' + unit['chapter_id'],
            'source_unit_ids': [unit['id']], 'event_ids': [event_id], 'cast': ['char-a'],
            'appearance_versions': {'char-a': 'base'},
            'action': unit['text'] + '：' + pid, 'shot': '中景', 'space': '人物位于测试房间中', 'expression': '平静',
            'state_before': {'char-a': copy.deepcopy(state)}, 'state_after': {'char-a': copy.deepcopy(state)},
            'dialogue': [{'kind': 'caption', 'text': unit['text']}]
        })
    page_idx = 1
    for chapter in source['chapters']:
        if not chapter['has_body']:
            continue
        ch_units = [u for u in body if u['chapter_id'] == chapter['id']]
        ch_pids = [f'p{body.index(u)+1}' for u in ch_units]
        i = 0
        while i < len(ch_pids):
            remaining = len(ch_pids) - i
            if remaining == 4:
                chunk = ch_pids[i:i+4]
                i += 4
            else:
                chunk = ch_pids[i:i+3]
                i += 3
            script['pages'].append({'id': f'page{page_idx}', 'chapter_id': chapter['id'], 'panel_ids': chunk, 'rows': [[p] for p in chunk], 'columns': 1})
            page_idx += 1
    return script


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='comic-tests-')
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / '中文作品'
        self.source = self.base / '原文.txt'
        self.source.write_text('第1章 来信\n甲拿起信封。\n甲拆开信封。\n甲取出信纸。\n第2章 回信\n甲铺开白纸。\n甲拿起毛笔。\n甲写下回答。\n', encoding='utf-8')
        self.invoke('init', source=[str(self.source)], title='通用测试')


    def invoke(self, name, **kwargs):
        from argparse import Namespace
        kwargs.setdefault('scope', None)
        kwargs.setdefault('bindings', None)
        kwargs.setdefault('outcome', 'failed')
        return cp.run(Namespace(command=name, project=str(self.root), **kwargs))


    def json_file(self, value, stem='data'):
        import uuid
        path = self.base / (stem + '-' + uuid.uuid4().hex[:8] + '.json')
        cp.atomic_json(path, value)
        return str(path)


    def prepare_script(self):
        self.invoke('confirm-source', note='已检查夹具来源与边界')
        project = cp.project_load(self.root)
        for chapter in project['source']['chapters']:
            if chapter['has_body']:
                self.invoke('mark-read', chapter=chapter['id'], note='已读取夹具正文：' + chapter['title'])
        script = fixture_script(project['source'])
        attach_adaptations(self.root, project['source'], script)
        self.invoke('set-script', file=self.json_file(script))
        return script


    def add_reviews(self, project_dir=None):
        target = Path(project_dir or self.root)
        project = cp.project_load(target)
        for kind, checks in cp.REVIEW_CHECKS.items():
            report = {'script_hash': cp.digest(project['script']),
                'reviewed_chapter_ids': [c['id'] for c in project['source']['chapters'] if c['has_body']],
                'checks': {k: True for k in checks}, 'evidence': '自动化夹具报告，只测试结构与关卡。', 'issues': []}
            from argparse import Namespace
            cp.run(Namespace(command='review', project=str(target), kind=kind, file=self.json_file(report),
                             scope=None, bindings=None, outcome='failed'))
        stages = cp.stage_batches(project)
        if stages:
            stage_dir = target / 'reports' / 'stage_reviews'
            stage_dir.mkdir(parents=True, exist_ok=True)
            for st in stages:
                st_hash = cp.stage_script_fingerprint(project, st['reviewed_chapters'])
                stage_report = {
                    'schema_version': cp.SCHEMA_VERSION,
                    'stage_range': {
                        'start_chapter_id': st['start_chapter_id'],
                        'end_chapter_id': st['end_chapter_id'],
                        'chapter_count': st['chapter_count'],
                        'stage_script_hash': st_hash,
                    },
                    'stage_script_hash': st_hash,
                    'reviewed_chapters': st['reviewed_chapters'],
                    'stage_approved': True,
                    'round_1_initial_audit': {
                        'initial_verdict': 'PASSED',
                        'auditors': [
                            {
                                'auditor_id': 'auditor_a',
                                'conversation_id': 'conv_fixture_a',
                                'scores': {'total_score': 95},
                                'deductions': [],
                                'issues': [],
                            },
                            {
                                'auditor_id': 'auditor_b',
                                'conversation_id': 'conv_fixture_b',
                                'scores': {'total_score': 92},
                                'deductions': [],
                                'issues': [],
                            },
                        ],
                    },
                    'round_2_verification': {
                        'stage_approved': True,
                        'approved_at': cp.now(),
                        'auditors_recheck': [
                            {
                                'auditor_id': 'auditor_a',
                                'final_scores': {'total_score': 95},
                                'verdict': 'PASSED',
                                'unresolved_issue_refs': [],
                            },
                            {
                                'auditor_id': 'auditor_b',
                                'final_scores': {'total_score': 92},
                                'verdict': 'PASSED',
                                'unresolved_issue_refs': [],
                            },
                        ],
                    },
                }
                cp.atomic_json(stage_dir / f"stage_{st['stage_index']:02d}.json", stage_report)


    def locked(self):
        self.prepare_script()
        self.add_reviews()
        self.invoke('lock-script')


    def image_file(self, color='white'):
        from PIL import Image
        import uuid
        path = self.base / ('mechanical-' + uuid.uuid4().hex[:8] + '.png')
        Image.new('RGB', (1080, 2400), color).save(path)
        return str(path)


    def qa(self, keys, reviewed_ids=None, image_sha256=None, reference_visual_key=None,
           attempt_bindings=None):
        report = {'checks': {k: True for k in keys},
                  'reviewed_ids': reviewed_ids or (['char-a'] if keys == cp.REFERENCE_CHECKS else [p['id'] for p in cp.project_load(self.root)['script']['panels']]),
                  'comparisons': [], 'findings': [],
                  'detail_notes': '机械夹具，只检查原生像素与报告字段，不宣称艺术细节达标。',
                  'elegance_notes': {k: '机械夹具，仅检查字段关卡，非真实美感结论。'
                                     for k in ('linework', 'color_and_light', 'visual_hierarchy')},
                  'evidence': '机械夹具，仅用于文件/流程回归，非真实图像验收。'}
        if image_sha256 is not None:
            report['image_sha256'] = image_sha256
        if reference_visual_key is not None:
            report['reference_visual_key'] = reference_visual_key
        if attempt_bindings is not None:
            report['attempt_bindings'] = attempt_bindings
        return self.json_file(report)


    def reference_qa(self, image_path, characters=None, subjects=None, purpose='combined'):
        if subjects is None:
            characters = characters or ['char-a']
            subjects = [{'character_id': cid, 'version_id': 'base', 'region': None} for cid in characters]
        inputs = self.invoke('qa-inputs', file=str(image_path),
                             bindings=self.json_file({'purpose': purpose, 'subjects': subjects}))
        return self.qa(cp.REFERENCE_CHECKS, inputs['reviewed_ids'], inputs['image_sha256'],
                       inputs['reference_visual_key'])


    def test_pre_art_gate_rejects_empty_and_partial_script(self):
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')
        script = self.prepare_script()
        script['panels'].pop()
        script['pages'].pop()
        self.invoke('set-script', file=self.json_file(script))
        self.assertTrue(any('Unmapped' in x or 'Chapter has no' in x for x in self.invoke('check-script')['errors']))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')


    def test_three_current_reviews_required(self):
        self.prepare_script()
        project = cp.project_load(self.root)
        for kind in ('coverage', 'continuity'):
            report = {'script_hash': cp.digest(project['script']), 'reviewed_chapter_ids': ['ch000001', 'ch000002'],
                      'checks': {k: True for k in cp.REVIEW_CHECKS[kind]}, 'evidence': '机械审查测试'}
            self.invoke('review', kind=kind, file=self.json_file(report))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')
        self.add_reviews()
        self.invoke('lock-script')
        self.assertTrue(self.invoke('assert-art')['allowed'])
        self.assertTrue((self.root / 'full-script.md').is_file())


    def test_source_modified_invalidates_lock(self):
        self.locked()
        self.source.write_text(self.source.read_text(encoding='utf-8') + '新增事件。', encoding='utf-8')
        self.assertTrue(self.invoke('status')['script_locked'])
        self.assertTrue(self.invoke('status')['external_source_warnings'])
        project = cp.project_load(self.root)
        archived = cp.inside(self.root, project['source']['files'][0]['archive_path'])
        archived.write_text(archived.read_text(encoding='utf-8') + '新增事件。', encoding='utf-8')
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')
        self.assertFalse(self.invoke('status')['script_locked'])


    def test_index_tampering_rejected(self):
        self.locked()
        project = cp.project_load(self.root)
        project['source']['units'][1]['text'] = '偷偷删改正文'
        cp.save(self.root, project)
        with self.assertRaises(cp.GateError):
            self.invoke('assert-art')


    def test_script_change_invalidates_old_reviews(self):
        self.locked()
        script = cp.project_load(self.root)['script']
        script['panels'][0]['action'] += '（新镜头）'
        self.invoke('set-script', file=self.json_file(script))
        with self.assertRaises(cp.GateError):
            self.invoke('lock-script')


    def test_duplicate_ids_and_wrong_speaker_rejected(self):
        script = self.prepare_script()
        script['panels'][1]['id'] = script['panels'][0]['id']
        script['panels'][0]['dialogue'] = [{'kind': 'speech', 'speaker': 'unknown', 'text': '你好'}]
        self.invoke('set-script', file=self.json_file(script))
        errors = self.invoke('check-script')['errors']
        self.assertTrue(any('duplicate' in x for x in errors))
        self.assertTrue(any('speaker' in x for x in errors))


    def test_state_transition_needs_source_evidence(self):
        script = self.prepare_script()
        for p in script['panels'][1:]:
            p['state_before']['char-a']['costume'] = 'coat-b'
            p['state_after']['char-a']['costume'] = 'coat-b'
        self.invoke('set-script', file=self.json_file(script))
        self.assertTrue(any('state' in x and ('change' in x or 'transition' in x) for x in self.invoke('check-script')['errors']))
        script['panels'][1]['state_transitions'] = [{'character_id': 'char-a', 'fields': ['costume'],
                                                  'reason': '夹具换装', 'source_unit_ids': script['panels'][1]['source_unit_ids']}]
        self.invoke('set-script', file=self.json_file(script))
        self.assertEqual([], self.invoke('check-script')['errors'])


    def test_major_unresolved_review_rejected(self):
        self.prepare_script()
        project = cp.project_load(self.root)
        report = {'script_hash': cp.digest(project['script']), 'reviewed_chapter_ids': ['ch000001', 'ch000002'],
                  'checks': {k: True for k in cp.REVIEW_CHECKS['coverage']}, 'evidence': '发现遗漏',
                  'issues': [{'severity': 'major', 'resolved': False}]}
        with self.assertRaises(cp.GateError):
            self.invoke('review', kind='coverage', file=self.json_file(report))


    def test_path_escape_rejected(self):
        with self.assertRaises(cp.GateError):
            cp.inside(self.root, '../outside.png')


    def test_unicode_encodings_duplicate_headings_and_empty_chapter(self):
        for encoding in ('utf-16', 'gb18030', 'utf-8-sig'):
            file = self.base / (encoding + '.txt')
            file.write_bytes('前言\n说明文字。\n第1章 同号\n正文。\n第1章 同号\n下一段。\n第2章 空章\n'.encode(encoding))
            source = cs.extract([file])
            self.assertEqual(4, len(source['chapters']))
            self.assertEqual(4, len({c['id'] for c in source['chapters']}))
            self.assertFalse(source['chapters'][-1]['has_body'])
            self.assertEqual(1, len(source['issues']))


    def test_docx_table_and_epub_spine_order(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('第1章 正文')
        doc.add_paragraph('文本。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表内书信。'
        path = self.base / '正文.docx'
        doc.save(path)
        self.assertIn('表内书信。', [u['text'] for u in cs.extract([path])['units']])
        path = self.base / '正文.epub'
        with zipfile.ZipFile(path, 'w') as book:
            book.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
            book.writestr('OPS/book.opf', '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
                '<item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/>'
                '</manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
            book.writestr('OPS/a.xhtml', '<html><body><p>后段。</p></body></html>')
            book.writestr('OPS/b.xhtml', '<html><body><p>先段。</p><script>隐藏内容</script></body></html>')
        self.assertEqual(['先段。', '后段。'], [u['text'] for u in cs.extract([path])['units']])


    def test_pdf_unreadable_page_flagged(self):
        from reportlab.pdfgen import canvas
        path = self.base / '空白页.pdf'
        pdf = canvas.Canvas(str(path))
        pdf.drawString(30, 700, 'Chapter 1')
        pdf.drawString(30, 680, 'Story text.')
        pdf.showPage()
        pdf.showPage()
        pdf.save()
        source = cs.extract([path])
        self.assertTrue(any(i.get('locator', {}).get('page') == 2 for i in source['issues']))


    def test_book_and_volume_hierarchy(self):
        book_dir = self.base / '神作小说'
        vol1_dir = book_dir / '第1卷'
        from argparse import Namespace
        init_res = cp.run(Namespace(command='init', project=str(vol1_dir),
                                    source=[str(self.source)], title='神作小说', volume='第1卷'))
        self.assertEqual(init_res['title'], '神作小说')
        self.assertEqual(init_res['volume'], '第1卷')
        loaded = cp.project_load(vol1_dir)
        self.assertEqual(loaded['title'], '神作小说')
        self.assertEqual(loaded['volume'], '第1卷')

        status_res = cp.run(Namespace(command='status', project=str(vol1_dir), plan=None))
        self.assertEqual(status_res['title'], '神作小说')
        self.assertEqual(status_res['volume'], '第1卷')

        # 校验剧本锁定后的 markdown 标题
        cp.run(Namespace(command='confirm-source', project=str(vol1_dir), note='来源确认', scope=None))
        for ch in loaded['source']['chapters']:
            if ch['has_body']:
                cp.run(Namespace(command='mark-read', project=str(vol1_dir), chapter=ch['id'], note='读完'))
        script = fixture_script(loaded['source'])
        attach_adaptations(vol1_dir, loaded['source'], script)
        script_file = self.json_file(script)
        cp.run(Namespace(command='set-script', project=str(vol1_dir), file=script_file))
        self.add_reviews(vol1_dir)
        cp.run(Namespace(command='lock-script', project=str(vol1_dir)))
        full_md = (vol1_dir / 'full-script.md').read_text(encoding='utf-8')
        self.assertIn('# 神作小说 · 第1卷 · 本卷漫画分镜剧本', full_md)


    def test_init_book_and_split_source(self):
        from argparse import Namespace
        book_dir = self.base / '星辰变漫改'
        raw_novel = self.base / '星辰变全书.txt'
        raw_novel.write_text(
            "第1卷 潜龙在渊\n第1章 异宝流星泪\n秦羽仰望星空。\n第2章 苦修\n少年挥汗如雨。\n"
            "第2卷 暴乱星海\n第1章 初入凶域\n波涛汹涌，海怪盘踞。\n",
            encoding='utf-8'
        )

        # 1. 运行 init-book，验证默认 action='copy' 保留原稿并归档副本至 source_texts
        init_res = cp.run(Namespace(
            command='init-book',
            book_dir=str(book_dir),
            title='星辰变',
            source=[str(raw_novel)],
            action='copy',
            description='热血修真史诗'
        ))
        self.assertTrue(init_res['ok'])
        self.assertEqual(init_res['title'], '星辰变')
        self.assertTrue(raw_novel.exists(), "默认 copy 模式下原小说文件必须予以保留")
        archived_novel = book_dir / 'source_texts' / '星辰变全书.txt'
        self.assertTrue(archived_novel.is_file(), "原稿副本应存在于 source_texts/")
        self.assertEqual(cp.sha_file(raw_novel), cp.sha_file(archived_novel), "归档副本哈希必须与原稿一致")
        self.assertTrue((book_dir / 'split_texts').is_dir(), "split_texts/ 目录应已创建")
        self.assertTrue((book_dir / 'docs').is_dir(), "docs/ 目录应已创建")

        # 验证显式 action='move' 时的移动行为
        move_source = self.base / '额外附录.txt'
        move_source.write_text("附录设定内容\n", encoding='utf-8')
        init_move_res = cp.run(Namespace(
            command='init-book',
            book_dir=str(book_dir),
            title='星辰变',
            source=[str(move_source)],
            action='move',
            description='热血修真史诗'
        ))
        self.assertTrue(init_move_res['ok'])
        self.assertFalse(move_source.exists(), "显式 move 模式下原文件应已被移动")
        self.assertTrue((book_dir / 'source_texts' / '额外附录.txt').is_file())
        for doc_name in ('overview.md', 'structure.md', 'worldview.md', 'characters.md', 'art_direction.md', 'progress.md'):
            self.assertTrue((book_dir / 'docs' / doc_name).is_file(), f"docs/{doc_name} 应存在")

        readme_text = (book_dir / 'README.md').read_text(encoding='utf-8')
        self.assertIn("# 《星辰变》漫画项目", readme_text)
        self.assertIn("全书分卷制作总览看板", readme_text)
        self.assertIn("怎么做：漫改全流程标准作业程序 (SOP)", readme_text)
        self.assertIn("各种规范引导文件在哪", readme_text)
        self.assertIn("原著真实性绝对铁律", readme_text)
        self.assertIn("docs/overview.md", readme_text)
        self.assertIn("docs/structure.md", readme_text)
        self.assertIn("docs/worldview.md", readme_text)
        self.assertIn("docs/characters.md", readme_text)
        self.assertIn("docs/art_direction.md", readme_text)
        self.assertIn("docs/progress.md", readme_text)
        self.assertIn("source_texts/", readme_text)
        self.assertIn("split_texts/", readme_text)
        self.assertIn("原著真实性原则", (book_dir / 'docs' / 'overview.md').read_text(encoding='utf-8'))

        # 2. 运行 split-source，验证切分结果统一存入 split_texts/
        split_res = cp.run(Namespace(
            command='split-source',
            book_dir=str(book_dir),
            file='星辰变全书.txt',
            output_dir=None,
            pattern=None
        ))
        self.assertTrue(split_res['ok'])
        self.assertEqual(split_res['segments_count'], 2)
        vol1_file = book_dir / 'split_texts' / 'vol_001_第1卷 潜龙在渊.txt'
        vol2_file = book_dir / 'split_texts' / 'vol_002_第2卷 暴乱星海.txt'
        self.assertTrue(vol1_file.is_file())
        self.assertTrue(vol2_file.is_file())
        self.assertTrue((book_dir / 'split_texts' / 'split_manifest.json').is_file())
        self.assertIn("秦羽仰望星空", vol1_file.read_text(encoding='utf-8'))
        self.assertIn("海怪盘踞", vol2_file.read_text(encoding='utf-8'))

        # 3. 验证以第1卷切割文本初始化该卷独立制作工作区，并验证生成该卷 docs 模块与精简 README.md
        vol1_project = book_dir / '第1卷'
        vol1_init = cp.run(Namespace(
            command='init',
            project=str(vol1_project),
            source=[str(vol1_file)],
            title='星辰变',
            volume='第1卷'
        ))
        self.assertEqual(vol1_init['title'], '星辰变')
        self.assertEqual(vol1_init['volume'], '第1卷')
        self.assertTrue((vol1_project / 'project.json').is_file())

        # 验证分卷 docs 目录及 6 个模块化文件
        vol1_docs = vol1_project / 'docs'
        self.assertTrue(vol1_docs.is_dir(), "分卷目录下应建立 docs/ 模块文件夹")
        for doc_file in ('info.md', 'status.md', 'commands.md', 'structure.md', 'notes.md', 'deliverables.md'):
            self.assertTrue((vol1_docs / doc_file).is_file(), f"分卷 docs/{doc_file} 应存在")

        self.assertIn("《星辰变》· 第1卷 基本信息", (vol1_docs / 'info.md').read_text(encoding='utf-8'))
        self.assertIn("原著真实性原则", (vol1_docs / 'info.md').read_text(encoding='utf-8'))
        self.assertIn("制作状态看板", (vol1_docs / 'status.md').read_text(encoding='utf-8'))
        self.assertIn("常用操作命令速查", (vol1_docs / 'commands.md').read_text(encoding='utf-8'))
        self.assertIn("工程目录结构", (vol1_docs / 'structure.md').read_text(encoding='utf-8'))
        self.assertIn("关键注记与特别要求", (vol1_docs / 'notes.md').read_text(encoding='utf-8'))
        self.assertIn("导出品交付路径", (vol1_docs / 'deliverables.md').read_text(encoding='utf-8'))

        # 验证精简的卷根目录 README.md（作为索引与导航）
        vol1_readme = vol1_project / 'README.md'
        self.assertTrue(vol1_readme.is_file(), "分卷目录下应自动生成 README.md")
        vol1_readme_text = vol1_readme.read_text(encoding='utf-8')
        self.assertIn("《星辰变》· 第1卷 漫画制作工程", vol1_readme_text)
        self.assertIn("原著真实性绝对原则", vol1_readme_text)
        self.assertIn("本卷关键状态速览", vol1_readme_text)
        self.assertIn("本卷模块化文档导航", vol1_readme_text)
        self.assertIn("docs/info.md", vol1_readme_text)
        self.assertIn("docs/status.md", vol1_readme_text)
        self.assertIn("docs/commands.md", vol1_readme_text)
        self.assertIn("docs/structure.md", vol1_readme_text)
        self.assertIn("docs/notes.md", vol1_readme_text)
        self.assertIn("docs/deliverables.md", vol1_readme_text)


    def reference(self):
        image=self.image_file()
        self.invoke('register-reference',characters=['char-a'],file=image,qa=self.reference_qa(image))

    def page_qa(self,page_id,attempt,image):
        inputs=self.invoke('qa-inputs',page=page_id,attempt=attempt,file=image)
        report=cp.load_json(self.qa(cp.PAGE_ART_CHECKS,inputs['reviewed_ids'],inputs['image_sha256']))
        report.update(attempt_binding=inputs['attempt_binding'],reviewed_page_ids=[page_id],
                      phone_reading_notes=[{'page_id':page_id,'preview_widths':[360,390,430],'min_body_css_px':18,
                      'evidence':'机械字段夹具，不宣称已读图或测字。'}])
        return self.json_file(report)

    def accept_all(self):
        self.reference()
        for page in cp.project_load(self.root)['script']['pages']:
            prompt=self.base/(page['id']+'.txt')
            prompt.write_text(self.invoke('build-prompt',page=page['id'],output=None)['prompt'],encoding='utf-8')
            result=self.invoke('begin-page',page=page['id'],prompt=str(prompt))
            if not result['generation_required']:continue
            image=self.image_file()
            self.invoke('finish-page',page=page['id'],attempt=result['attempt'],file=image,
                        qa=self.page_qa(page['id'],result['attempt'],image))

    def review_and_export(self):
        layout=cp.project_load(self.root)['layout']
        report={'input_hash':layout['input_hash'],'reviewed_page_ids':[p['id'] for p in layout['pages']],
                'checks':{k:True for k in cp.LAYOUT_CHECKS},'evidence':'机械字段夹具，不声明艺术可读性。',
                'elegance_notes':{k:'机械字段夹具' for k in ('linework','color_and_light','visual_hierarchy')},
                'phone_reading_notes':[{'page_id':p['id'],'preview_widths':[360,390,430],'min_body_css_px':18,
                                       'evidence':'机械字段夹具'} for p in layout['pages']]}
        self.invoke('review-layout',file=self.json_file(report))
        self.invoke('export')

    def test_cli_smoke_build_prompt(self):
        import subprocess
        self.locked()
        self.reference()
        script = cp.project_load(self.root)['script']
        page_id = script['pages'][0]['id']

        # Test 1: build-prompt with --page produces JSON output with prompt
        cmd = [sys.executable, str(Path(cp.__file__)), 'build-prompt',
               '--project', str(self.root), '--page', page_id]
        res = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(0, res.returncode, f"build-prompt failed: {res.stderr}")
        out = json.loads(res.stdout)
        self.assertTrue(out.get('ok'))
        self.assertIn('prompt', out)
        self.assertEqual(page_id, out['page_id'])

        # Test 2: build-prompt with --notes and --output file
        out_prompt_file = self.base / 'test_smoke_prompt.txt'
        cmd_notes = [sys.executable, str(Path(cp.__file__)), 'build-prompt',
                     '--project', str(self.root), '--page', page_id,
                     '--notes', '重要镜头提示', '--output', str(out_prompt_file)]
        res_notes = subprocess.run(cmd_notes, capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(0, res_notes.returncode, f"build-prompt with notes failed: {res_notes.stderr}")
        self.assertTrue(out_prompt_file.is_file())
        self.assertIn('重要镜头提示', out_prompt_file.read_text(encoding='utf-8'))

        # Test 3: verify abbreviation --pag is rejected because allow_abbrev=False
        cmd_abbrev = [sys.executable, str(Path(cp.__file__)), 'build-prompt',
                      '--project', str(self.root), '--pag', page_id]
        res_abbrev = subprocess.run(cmd_abbrev, capture_output=True, text=True, encoding='utf-8')
        self.assertNotEqual(0, res_abbrev.returncode, "Abbreviated flag --pag must be rejected")

    def test_stage_review_bypass_blocked_and_fingerprint_stale(self):
        p_dir = self.base / '五章项目'
        src = self.base / '五章原文.txt'
        src_text = "".join(f"第{i}章 章节{i}\n人物甲动作A{i}。\n人物甲动作B{i}。\n人物甲动作C{i}。\n" for i in range(1, 6))
        src.write_text(src_text, encoding='utf-8')
        from argparse import Namespace
        cp.run(Namespace(command='init', project=str(p_dir), source=[str(src)], title='五章工程', volume='第1卷'))
        cp.run(Namespace(command='confirm-source', project=str(p_dir), note='来源核验', scope=None))
        proj = cp.project_load(p_dir)
        for ch in proj['source']['chapters']:
            if ch['has_body']:
                cp.run(Namespace(command='mark-read', project=str(p_dir), chapter=ch['id'], note='已读'))
        script = fixture_script(proj['source'])
        attach_adaptations(p_dir, proj['source'], script)
        cp.run(Namespace(command='set-script', project=str(p_dir), file=self.json_file(script)))

        # Add volume reviews but NO stage review
        proj = cp.project_load(p_dir)
        for kind, checks in cp.REVIEW_CHECKS.items():
            rep = {'script_hash': cp.digest(proj['script']),
                   'reviewed_chapter_ids': [c['id'] for c in proj['source']['chapters'] if c['has_body']],
                   'checks': {k: True for k in checks}, 'evidence': '测试证据', 'issues': []}
            cp.run(Namespace(command='review', project=str(p_dir), kind=kind, file=self.json_file(rep)))

        # 1. check-stage-review must report missing stage review
        with self.assertRaises(cp.GateError) as ctx_stage:
            cp.run(Namespace(command='check-stage-review', project=str(p_dir), file=None))
        self.assertIn('No reports/stage_reviews directory found', str(ctx_stage.exception))

        # 2. lock-script must fail because stage review is missing
        with self.assertRaises(cp.GateError) as ctx:
            cp.run(Namespace(command='lock-script', project=str(p_dir)))
        self.assertIn('Missing required stage reviews', str(ctx.exception))

        # 3. Add valid stage review
        self.add_reviews(p_dir)
        check_ok = cp.run(Namespace(command='check-stage-review', project=str(p_dir), file=None))
        self.assertTrue(check_ok['ok'])
        lock_res = cp.run(Namespace(command='lock-script', project=str(p_dir)))
        self.assertTrue(lock_res['ok'])

        # 4. Modify chapter 1 panels in script -> stage review must become stale
        script['panels'][0]['action'] += '（新改动动作）'
        cp.run(Namespace(command='set-script', project=str(p_dir), file=self.json_file(script)))
        with self.assertRaises(cp.GateError) as ctx_stale:
            cp.run(Namespace(command='check-stage-review', project=str(p_dir), file=None))
        self.assertTrue(any(word in str(ctx_stale.exception) for word in ('mismatch', 'changed', 'stale')))
        with self.assertRaises(cp.GateError) as ctx2:
            cp.run(Namespace(command='lock-script', project=str(p_dir)))
        self.assertTrue(any(word in str(ctx2.exception) for word in ('mismatch', 'changed', 'stale')))

    def test_preview_page_generation(self):
        self.locked()
        script = cp.project_load(self.root)['script']
        page_id = script['pages'][0]['id']
        img = self.image_file(color='blue')
        prev_res = self.invoke('preview-page', page=page_id, file=img)
        self.assertTrue(prev_res['ok'])
        self.assertEqual(page_id, prev_res['page_id'])
        for w in (360, 390, 430):
            pfile = self.root / 'previews' / 'pages' / page_id / f'{w}.png'
            self.assertTrue(pfile.is_file(), f"Preview file {pfile} should exist")
            from PIL import Image
            with Image.open(pfile) as pimg:
                self.assertEqual(w, pimg.width)

    def test_cost_estimation_and_typeset_preflight(self):
        self.locked()
        cost_res = self.invoke('estimate-cost')
        self.assertTrue(cost_res['ok'])
        self.assertIn('estimated_pages', cost_res)
        self.assertIn('estimated_art_calls', cost_res)
        self.assertIn('cost_visibility', cost_res)

        typeset_ok = self.invoke('preflight-typeset')
        self.assertTrue(typeset_ok['ok'])
        self.assertEqual(0, len(typeset_ok['errors']))

        # Catch dialogue overflow
        script = cp.project_load(self.root)['script']
        huge_text = '这是一段极长的测试对白，用来检验排字容量检查。' * 500
        script['panels'][0]['dialogue'].append({'kind': 'speech', 'speaker': 'char-a', 'text': huge_text})
        self.invoke('set-script', file=self.json_file(script))
        typeset_fail = self.invoke('preflight-typeset')
        self.assertFalse(typeset_fail['ok'])
        self.assertTrue(any('exceeds' in e and 'capacity' in e for e in typeset_fail['errors']))

    def exported(self):
        self.locked()
        self.accept_all()
        self.invoke('prepare-pages')
        self.review_and_export()


if __name__=='__main__':
    unittest.main()
